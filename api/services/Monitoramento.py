import os
from datetime import datetime
from openpyxl import Workbook, load_workbook
import psutil

class Monitoramento:
    def __init__(self):
        # Nome do arquivo baseado na data de hoje
        self.data_atual = datetime.now().strftime('%Y-%m-%d')
        self.arquivo = f"./desempenho/{self.data_atual}.xlsx"
        # Verifica se o arquivo já existe
        if not os.path.exists(self.arquivo):
            self.criar_planilha()
        else:
            self.planilha = load_workbook(self.arquivo)
            self.sheet = self.planilha.active

    def criar_planilha(self):
        self.planilha = Workbook()
        self.sheet = self.planilha.active
        # Cabeçalhos
        self.sheet.append(['Data', 'Hora', 'PID', 'IP', 'Memória (MB)', 'Ação'])
        self.planilha.save(self.arquivo)

    def registrar_acao(self, acao, ip=''):
        # Atualiza o arquivo sempre com a data mais atual
        self.data_atual = datetime.now().strftime('%Y-%m-%d')
        self.arquivo = f"./desempenho/{self.data_atual}.xlsx"

        # Verifica se o arquivo existe a cada registro, se necessário reabre
        if not os.path.exists(self.arquivo):
            self.criar_planilha()
        else:
            self.planilha = load_workbook(self.arquivo)
            self.sheet = self.planilha.active

        agora = datetime.now()
        data = agora.strftime('%Y-%m-%d')
        hora = agora.strftime('%H:%M:%S')
        pid = os.getpid()

        # Pega a memória usada pelo processo atual
        memoria = psutil.Process(pid).memory_info().rss / (1024 * 1024)  # MB

        # Adiciona a linha na planilha
        self.sheet.append([data, hora, pid, ip, memoria, acao])
        self.planilha.save(self.arquivo)